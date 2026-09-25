"""Personal notes workspace built with Flask."""
import hmac
import io
import json
import os
import re
import secrets
from datetime import datetime, timedelta
from functools import wraps
from pathlib import Path

from flask import Flask, abort, flash, redirect, render_template, request, send_file, session, url_for
from sqlalchemy import inspect, or_, text
from werkzeug.security import check_password_hash, generate_password_hash

from models import Note, User, db

ROOT = Path(__file__).resolve().parent
app = Flask(__name__)
app.instance_path = str(ROOT / "instance")
Path(app.instance_path).mkdir(exist_ok=True)
secret_file = Path(app.instance_path) / "secret_key"
if not os.environ.get("SECRET_KEY") and not secret_file.exists():
    secret_file.write_text(secrets.token_hex(32), encoding="utf-8")
app.config.update(
    SECRET_KEY=os.environ.get("SECRET_KEY") or secret_file.read_text(encoding="utf-8"),
    SQLALCHEMY_DATABASE_URI=os.environ.get("DATABASE_URL") or "sqlite:///notes.db",
    SQLALCHEMY_TRACK_MODIFICATIONS=False,
    MAX_CONTENT_LENGTH=3 * 1024 * 1024,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
)
db.init_app(app)


def migrate_database():
    """Preserve original SQLite notes while adding workspace fields."""
    with app.app_context():
        db.create_all()
        columns = {column["name"] for column in inspect(db.engine).get_columns("note")}
        additions = {"updated_at": "DATETIME", "tags": "VARCHAR(500) DEFAULT ''", "is_pinned": "BOOLEAN DEFAULT 0", "is_archived": "BOOLEAN DEFAULT 0"}
        for name, definition in additions.items():
            if name not in columns:
                db.session.execute(text(f"ALTER TABLE note ADD COLUMN {name} {definition}"))
        db.session.commit()


migrate_database()


@app.before_request
def protect_forms():
    if request.method in {"POST", "PUT", "PATCH", "DELETE"}:
        expected = session.get("csrf_token", "")
        supplied = request.form.get("csrf_token", "")
        if not expected or not hmac.compare_digest(expected, supplied):
            abort(400, "表单已过期，请刷新页面后重试。")


@app.context_processor
def template_globals():
    if "csrf_token" not in session:
        session["csrf_token"] = secrets.token_urlsafe(32)
    return {"csrf_token": session["csrf_token"], "clear_draft_key": session.pop("clear_draft_key", "")}


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if "user_id" not in session:
            return redirect(url_for("login"))
        return view(*args, **kwargs)
    return wrapped


def admin_required(view):
    @wraps(view)
    @login_required
    def wrapped(*args, **kwargs):
        if session.get("role") != "admin":
            abort(403)
        return view(*args, **kwargs)
    return wrapped


def own_note(note_id):
    note = db.session.get(Note, note_id)
    if note is None:
        abort(404)
    if note.user_id != session["user_id"] and session.get("role") != "admin":
        abort(403)
    return note


def normalize_tags(raw):
    parts = re.split(r"[,，、\n]+", raw or "")
    return ",".join(dict.fromkeys(part.strip().lstrip("#")[:30] for part in parts if part.strip()))[:500]


def parse_day(value):
    try:
        return datetime.strptime(value, "%Y-%m-%d") if value else None
    except ValueError:
        abort(400, "日期格式无效。")


def filtered_notes(all_users=False, values=None):
    values = values or request.args
    query = Note.query if all_users else Note.query.filter_by(user_id=session["user_id"])
    keyword = (values.get("keyword") or "").strip()[:100]
    if keyword:
        term = f"%{keyword}%"
        query = query.filter(or_(Note.title.ilike(term), Note.content.ilike(term), Note.tags.ilike(term)))
    tag = (values.get("tag") or "").strip()[:30]
    if tag:
        query = query.filter(or_(Note.tags == tag, Note.tags.like(f"{tag},%"), Note.tags.like(f"%,{tag}"), Note.tags.like(f"%,{tag},%")))
    start, end = parse_day(values.get("start_date")), parse_day(values.get("end_date"))
    if start:
        query = query.filter(Note.create_time >= start)
    if end:
        query = query.filter(Note.create_time < end + timedelta(days=1))
    view = values.get("view", "active")
    if view == "archived":
        query = query.filter(Note.is_archived.is_(True))
    elif view != "all":
        query = query.filter(or_(Note.is_archived.is_(False), Note.is_archived.is_(None)))
    if view == "pinned":
        query = query.filter(Note.is_pinned.is_(True))
    return query.order_by(Note.is_pinned.desc(), Note.updated_at.desc(), Note.create_time.desc())


@app.route("/", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = (request.form.get("username") or "").strip()
        password = request.form.get("password") or ""
        user = User.query.filter_by(username=username).first()
        hashed = user and user.password.startswith(("scrypt:", "pbkdf2:"))
        valid = user and (check_password_hash(user.password, password) if hashed else hmac.compare_digest(user.password, password))
        if valid:
            if not hashed:
                user.password = generate_password_hash(password)
                db.session.commit()
            session.clear()
            session.update(user_id=user.id, username=user.username, role=user.role)
            flash("欢迎回来。", "success")
            return redirect(url_for("admin") if user.role == "admin" else url_for("home"))
        flash("账号或密码错误。", "error")
    elif "user_id" in session:
        return redirect(url_for("admin") if session.get("role") == "admin" else url_for("home"))
    return render_template("login.html")


@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        username = (request.form.get("username") or "").strip()
        password = request.form.get("password") or ""
        if not 3 <= len(username) <= 50 or not 8 <= len(password) <= 128:
            flash("用户名需 3–50 字，密码需 8–128 字。", "error")
        elif User.query.filter_by(username=username).first():
            flash("用户名已存在。", "error")
        else:
            db.session.add(User(username=username, password=generate_password_hash(password), role="user"))
            db.session.commit()
            flash("注册成功，请登录。", "success")
            return redirect(url_for("login"))
    return render_template("register.html")


@app.post("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/home")
@login_required
def home():
    if session.get("role") == "admin":
        return redirect(url_for("admin"))
    notes = filtered_notes().all()
    all_notes = Note.query.filter_by(user_id=session["user_id"]).all()
    tags = sorted({tag for note in all_notes for tag in note.tag_list}, key=str.casefold)
    return render_template("home.html", notes=notes, tags=tags, total=len(all_notes), active=sum(not n.is_archived for n in all_notes), pinned=sum(bool(n.is_pinned and not n.is_archived) for n in all_notes), archived=sum(bool(n.is_archived) for n in all_notes), view=request.args.get("view", "active"))


@app.route("/write", methods=["GET", "POST"])
@login_required
def write():
    if request.method == "POST":
        title = (request.form.get("title") or "").strip()[:200]
        if title:
            note = Note(title=title, content=(request.form.get("content") or "").strip(), user_id=session["user_id"], tags=normalize_tags(request.form.get("tags")), is_pinned=bool(request.form.get("is_pinned")), updated_at=datetime.now())
            db.session.add(note)
            db.session.commit()
            session["clear_draft_key"] = "new-note"
            flash("笔记已保存。", "success")
            return redirect(url_for("note_detail", id=note.id))
        flash("请填写标题。", "error")
    return render_template("editor.html", note=None, heading="新建笔记")


@app.route("/note/<int:id>")
@login_required
def note_detail(id):
    return render_template("note_detail.html", note=own_note(id))


@app.route("/edit/<int:id>", methods=["GET", "POST"])
@login_required
def edit(id):
    note = own_note(id)
    if request.method == "POST":
        title = (request.form.get("title") or "").strip()[:200]
        if title:
            note.title = title
            note.content = (request.form.get("content") or "").strip()
            note.tags = normalize_tags(request.form.get("tags"))
            note.is_pinned = bool(request.form.get("is_pinned"))
            note.updated_at = datetime.now()
            db.session.commit()
            session["clear_draft_key"] = f"note-{id}"
            flash("修改已保存。", "success")
            return redirect(url_for("note_detail", id=id))
        flash("请填写标题。", "error")
    return render_template("editor.html", note=note, heading="编辑笔记")


@app.post("/delete/<int:id>")
@login_required
def delete(id):
    db.session.delete(own_note(id))
    db.session.commit()
    flash("笔记已删除。", "success")
    return redirect(url_for("admin_notes") if session.get("role") == "admin" else url_for("home"))


@app.post("/note/<int:id>/toggle/<action>")
@login_required
def toggle_note(id, action):
    note = own_note(id)
    if action not in {"pin", "archive"}:
        abort(404)
    field = "is_pinned" if action == "pin" else "is_archived"
    setattr(note, field, not bool(getattr(note, field)))
    note.updated_at = datetime.now()
    db.session.commit()
    return redirect(request.referrer if request.referrer and request.referrer.startswith(request.host_url) else url_for("home"))


@app.route("/search", methods=["GET", "POST"])
@login_required
def search():
    if request.method == "POST":
        return redirect(url_for("search", **{key: value for key, value in request.form.items() if key != "csrf_token"}))
    return render_template("search.html", notes=filtered_notes(all_users=session.get("role") == "admin").all())


def export_bytes(notes, kind):
    if kind == "json":
        data = [{"title": n.title, "content": n.content, "tags": n.tag_list, "created_at": n.create_time.isoformat(), "updated_at": n.updated_at.isoformat() if n.updated_at else None, "pinned": bool(n.is_pinned), "archived": bool(n.is_archived)} for n in notes]
        return json.dumps({"version": 1, "notes": data}, ensure_ascii=False, indent=2).encode("utf-8"), "application/json", "json"
    if kind in {"txt", "md"}:
        blocks = [f"{'# ' if kind == 'md' else ''}{n.title}\n{n.create_time:%Y-%m-%d %H:%M} · {', '.join(n.tag_list)}\n\n{n.content}" for n in notes]
        return ("\n\n---\n\n".join(blocks) + "\n").encode("utf-8-sig"), "text/plain; charset=utf-8", kind
    if kind == "word":
        from docx import Document
        document = Document()
        document.add_heading("我的学习笔记", 0)
        for n in notes:
            document.add_heading(n.title, 1)
            document.add_paragraph(n.content)
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"
    if kind == "pdf":
        from xml.sax.saxutils import escape
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        body = ParagraphStyle("body", fontName="STSong-Light", fontSize=11, leading=18, spaceAfter=14)
        title = ParagraphStyle("title", parent=body, fontSize=17, leading=24)
        story = []
        for n in notes:
            story.extend([Paragraph(escape(n.title), title), Paragraph(escape(n.content).replace("\n", "<br/>"), body), Spacer(1, 15)])
        buffer = io.BytesIO()
        SimpleDocTemplate(buffer).build(story)
        return buffer.getvalue(), "application/pdf", "pdf"
    abort(400, "不支持的导出格式。")


@app.route("/export", methods=["GET", "POST"])
@login_required
def export():
    all_users = session.get("role") == "admin"
    if request.method == "GET":
        return render_template("export.html", notes=filtered_notes(all_users=all_users, values={"view": "all"}).all())
    notes = filtered_notes(all_users=all_users, values={**request.form, "view": "all"})
    raw_ids = request.form.getlist("note_ids")
    if raw_ids:
        try:
            ids = [int(value) for value in raw_ids]
        except ValueError:
            abort(400)
        notes = notes.filter(Note.id.in_(ids))
    selected = notes.all()
    if not selected:
        flash("没有符合条件的笔记。", "error")
        return redirect(url_for("export"))
    payload, mime, extension = export_bytes(selected, request.form.get("format", "json"))
    return send_file(io.BytesIO(payload), mimetype=mime, as_attachment=True, download_name=f"notes-{datetime.now():%Y%m%d-%H%M}.{extension}")


@app.post("/import")
@login_required
def import_notes():
    upload = request.files.get("file")
    if not upload or not upload.filename.lower().endswith(".json"):
        flash("请选择本应用导出的 JSON 文件。", "error")
        return redirect(url_for("export"))
    try:
        items = json.load(upload)["notes"]
        if not isinstance(items, list) or len(items) > 500:
            raise ValueError
        prepared = []
        for item in items:
            if not isinstance(item, dict) or not isinstance(item.get("title"), str) or not item["title"].strip() or not isinstance(item.get("content", ""), str):
                raise ValueError
            tags = item.get("tags", [])
            created_at = datetime.fromisoformat(item["created_at"]) if item.get("created_at") else datetime.now()
            updated_at = datetime.fromisoformat(item["updated_at"]) if item.get("updated_at") else created_at
            prepared.append(Note(title=item["title"].strip()[:200], content=item.get("content", "")[:100000], tags=normalize_tags(",".join(tags) if isinstance(tags, list) and all(isinstance(tag, str) for tag in tags) else ""), is_pinned=bool(item.get("pinned")), is_archived=bool(item.get("archived")), user_id=session["user_id"], create_time=created_at, updated_at=updated_at))
    except (ValueError, TypeError, KeyError, UnicodeDecodeError, json.JSONDecodeError):
        flash("备份文件格式无效。", "error")
        return redirect(url_for("export"))
    db.session.add_all(prepared)
    db.session.commit()
    flash(f"已导入 {len(prepared)} 篇笔记。", "success")
    return redirect(url_for("home") if session.get("role") != "admin" else url_for("admin_notes"))


@app.route("/admin")
@admin_required
def admin():
    return render_template("admin.html", count=Note.query.count(), user_count=User.query.count(), notes=Note.query.order_by(Note.create_time.desc()).limit(5).all())


@app.route("/admin/notes")
@admin_required
def admin_notes():
    return render_template("admin_notes.html", notes=filtered_notes(all_users=True).all())


@app.route("/admin/note/<int:id>")
@admin_required
def admin_note_detail(id):
    return redirect(url_for("note_detail", id=id))


@app.route("/admin/edit/<int:id>", methods=["GET", "POST"])
@admin_required
def admin_edit(id):
    return edit(id)


@app.post("/admin/delete/<int:id>")
@admin_required
def admin_delete(id):
    return delete(id)


@app.errorhandler(400)
@app.errorhandler(403)
@app.errorhandler(404)
def error_page(error):
    return render_template("error.html", code=error.code, message=error.description), error.code


if __name__ == "__main__":
    app.run(debug=os.environ.get("FLASK_DEBUG") == "1")
